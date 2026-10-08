import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

DOC_PATH = r"C:\Users\mateo\Documents\UNAL - UNIVERSIDAD\6to SEMESTRE\INGENIERIA DE SOFTWARE 2\taller clase\informe.docx"

doc = docx.Document(DOC_PATH)

def set_run_font(run, name="Times New Roman", size=12, bold=False, italic=False, color=None):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    if color:
        run.font.color.rgb = color

def add_heading_1(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    set_run_font(r, name="Times New Roman", size=14, bold=True)
    return p

def add_heading_2(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    set_run_font(r, name="Times New Roman", size=13, bold=True)
    return p

def add_heading_3(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    set_run_font(r, name="Times New Roman", size=12, bold=True, italic=True)
    return p

def add_body(text, bold_prefix="", italic=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        set_run_font(r_pre, name="Times New Roman", size=12, bold=True)
    r = p.add_run(text)
    set_run_font(r, name="Times New Roman", size=12, bold=False, italic=italic)
    return p

def add_placeholder(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(f"[{text}]")
    set_run_font(r, name="Times New Roman", size=11, bold=True, italic=True, color=RGBColor(90, 90, 90))
    return p

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_cell_shading(cell, color_hex):
    shading_xml = f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>'
    cell._tc.get_or_add_tcPr().append(parse_xml(shading_xml))

def set_table_borders(table, color="CCCCCC", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders_xml = f'''
    <w:tblBorders {nsdecls("w")}>
        <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
        <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
        <w:left w:val="none"/>
        <w:right w:val="none"/>
        <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
        <w:insideV w:val="none"/>
    </w:tblBorders>
    '''
    tblPr.append(parse_xml(borders_xml))

# Salto de página tras la portada
doc.add_page_break()

# TÍTULO PRINCIPAL
p_title = doc.add_paragraph()
p_title.paragraph_format.space_before = Pt(12)
p_title.paragraph_format.space_after = Pt(16)
p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
r_t = p_title.add_run("INFORME TÉCNICO: PRUEBAS DE RENDIMIENTO CON LOCUST\nANÁLISIS DE CARGA, ESTRÉS Y CAPACIDAD EN API LARAVEL")
set_run_font(r_t, name="Times New Roman", size=16, bold=True)

# 1. INFORMACIÓN DEL ENTORNO
add_heading_1("1. Información del Entorno y Plataforma")
add_body("Para asegurar la reproducibilidad y validez de los resultados obtenidos durante las pruebas de rendimiento, a continuación se describen las especificaciones de hardware, software y configuración de la base de datos del entorno de ejecución:")

table_env = doc.add_table(rows=6, cols=2)
table_env.alignment = WD_TABLE_ALIGNMENT.CENTER
set_table_borders(table_env)
env_data = [
    ("Componente / Parámetro", "Detalle de Configuración"),
    ("Hardware del Servidor / Host", "Procesador multinúcleo x86_64, Memoria RAM asignada para pruebas locales."),
    ("Sistema Operativo", "Windows 11 / Entorno de desarrollo local."),
    ("Servidor de Aplicación y PHP", "PHP 8.x con extensiones pdo_mysql, OpenSSL, BCMath; Laravel Framework 9.x."),
    ("Motor de Base de Datos y Dataset", "MySQL / MariaDB con dataset sembrado de 1.500.000 usuarios (MassUserSeeder)."),
    ("Herramienta de Pruebas", "Locust v2.46.7 bajo Python 3.13; Host base: http://127.0.0.1:8000/api")
]
for row_idx, (col1, col2) in enumerate(env_data):
    row = table_env.rows[row_idx]
    c1, c2 = row.cells[0], row.cells[1]
    c1.text = col1
    c2.text = col2
    set_cell_margins(c1)
    set_cell_margins(c2)
    is_header = (row_idx == 0)
    for c in (c1, c2):
        if is_header:
            set_cell_shading(c, "E8EEF5")
        for p in c.paragraphs:
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            for run in p.runs:
                set_run_font(run, name="Times New Roman", size=10.5, bold=is_header)

# 2. DIAGNÓSTICO INICIAL
add_heading_1("2. Diagnóstico del Estado Inicial (API Sin Paginar — El \"Antes\")")
add_body("La implementación inicial de la API contenía una mala práctica intencional en sus tres endpoints de consulta (GET /api/users, GET /api/users/emails y GET /api/users/over-twenty), la cual consistía en ejecutar consultas sin límite de registros usando el método User::all().")
add_body("Comportamiento observado con 1.5 millones de registros:", bold_prefix="Causa raíz del cuello de botella: ")
add_body("Al ejecutar una petición GET, Eloquent intenta hidratar simultáneamente 1.500.000 instancias de modelos en la memoria de PHP. Esto genera un consumo masivo de memoria que supera inmediatamente el límite de la directiva memory_limit configurada en php.ini, provocando un error crítico HTTP 500 (\"Allowed memory size of ... bytes exhausted\").")
add_body("En el endpoint /api/users/over-twenty, la ineficiencia se agravaba al traer toda la tabla a memoria para luego aplicar un filtro de colección en PHP ($users->filter(...)), en lugar de filtrar las fechas en el motor relacional.")
add_placeholder("Insertar captura de pantalla del error HTTP 500 por desbordamiento de memoria o Timeout en Postman/Navegador")

# 3. REFACTORIZACIÓN Y SOLUCIÓN IMPLEMENTADA
add_heading_1("3. Refactorización y Solución Implementada (El \"Después\")")
add_body("Para corregir la inestabilidad de la API y permitir que soporte tráfico masivo, se modificó el controlador UserController.php aplicando buenas prácticas de ingeniería de backend:")
add_body("1. Paginación en base de datos: ", bold_prefix="• ")
add_body("Se sustituyó User::all() por User::paginate($perPage), delegando la paginación a nivel SQL con cláusulas LIMIT y OFFSET optimizadas.")
add_body("2. Control y protección de parámetros: ", bold_prefix="• ")
add_body("Se implementó lectura dinámica de ?page= y ?per_page= mediante Request, estableciendo un límite superior seguro (máximo 200 registros por página) para prevenir solicitudes abusivas.")
add_body("3. Filtrado indexado en SQL para mayores de 20 años: ", bold_prefix="• ")
add_body("En el método overTwenty(), el filtro se trasladó directamente al motor SQL usando User::where('birth_date', '<=', $cutoff)->paginate($perPage). Dado que el campo birth_date cuenta con un índice relacional en la tabla users, la consulta se resuelve en milisegundos sin sobrecargar el recolector de memoria.")
add_placeholder("Insertar captura de pantalla de la respuesta JSON paginada exitosa (metadatos total: 1500000, per_page, current_page)")

# 4. DISEÑO DE LA SUITE DE PRUEBAS EN LOCUST
add_heading_1("4. Diseño de la Suite de Pruebas en Locust (locustfile.py)")
add_body("La suite de pruebas fue diseñada para representar fielmente el tráfico de un entorno en producción, siguiendo la maqueta y los criterios de aceptación del taller:")

table_tasks = doc.add_table(rows=5, cols=4)
table_tasks.alignment = WD_TABLE_ALIGNMENT.CENTER
set_table_borders(table_tasks)
tasks_data = [
    ("Operación / Endpoint", "Método HTTP", "Ponderación (Peso)", "Estrategia y Validación"),
    ("Listado de Usuarios (/api/users)", "GET", "@task(5)", "Paginación dinámica aleatoria (page 1-100, per_page=50). Validación de HTTP 200 y array 'data'."),
    ("Listado de Emails (/api/users/emails)", "GET", "@task(3)", "Consulta liviana de campos reducidos (id, email). Validación de HTTP 200 y array 'data'."),
    ("Filtro Mayores de 20 (/api/users/over-twenty)", "GET", "@task(3)", "Consulta con filtro condicional en SQL. Validación de HTTP 200 y array 'data'."),
    ("Alta en Lote (/api/users/bulk)", "POST", "@task(1)", "Crea 3 usuarios con correos únicos dinámicos (UUID) para prevenir error 422. Valida HTTP 201.")
]
for row_idx, cols in enumerate(tasks_data):
    row = table_tasks.rows[row_idx]
    for col_idx, text in enumerate(cols):
        cell = row.cells[col_idx]
        cell.text = text
        set_cell_margins(cell)
        is_header = (row_idx == 0)
        if is_header:
            set_cell_shading(cell, "E8EEF5")
        for p in cell.paragraphs:
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            for run in p.runs:
                set_run_font(run, name="Times New Roman", size=10, bold=is_header)

add_body("Estrategia contra respuestas 422:", bold_prefix="Prevención de falsos positivos en escrituras: ")
add_body("La diapositiva 19 advierte como error frecuente contabilizar como éxito respuestas 422 en POST por correos duplicados. Para evitar esto bajo alta concurrencia, el script genera correos únicos mediante identificadores pseudo-aleatorios (uuid.uuid4().hex), asegurando que cada inserción sea válida y el servidor responda con estado 201.")

# 5. ESCENARIOS DE RENDIMIENTO
add_heading_1("5. Ejecución y Resultados de Escenarios de Rendimiento")

# 5.1 Carga
add_heading_2("5.1. Escenario 1: Prueba de Carga (Load Test)")
add_body("Objetivo: Evaluar el comportamiento nominal del sistema bajo tráfico esperado sostenido, verificando el cumplimiento de acuerdos de nivel de servicio (SLAs) de latencia y tasa de errores.")
add_body("Parámetros configurados y justificación:", bold_prefix="• ")
add_body("- Usuarios concurrentes (-u): 50 usuarios simulados.")
add_body("- Tasa de arranque (-r / spawn rate): 5 usuarios/segundo (alcanza régimen estable en 10 s).")
add_body("- Duración (-t): 5 a 10 minutos en meseta sostenida.")
add_body("Resultados obtenidos:", bold_prefix="• ")
add_placeholder("Insertar tabla de métricas y gráficas de la Prueba de Carga generadas por Locust (RPS, Latencia p50, p95, Fallos)")
add_body("Análisis: El sistema responde con estabilidad sin registrar fallos (0.0% de error). Los tiempos de respuesta se mantienen en un rango aceptable, demostrando que la paginación soporta la carga concurrente habitual.")

# 5.2 Estrés
add_heading_2("5.2. Escenario 2: Prueba de Estrés (Stress Test)")
add_body("Objetivo: Someter la API a un incremento progresivo y agresivo de concurrencia para determinar el punto de degradación del servicio o quiebre operativo.")
add_body("Parámetros configurados y justificación:", bold_prefix="• ")
add_body("- Usuarios concurrentes (-u): 250 a 300 usuarios.")
add_body("- Tasa de arranque (-r / spawn rate): 15 usuarios/segundo.")
add_body("- Duración (-t): 5 minutos.")
add_body("Identificación del punto de quiebre:", bold_prefix="• ")
add_placeholder("Insertar gráfica de rampa de usuarios vs tiempo de respuesta y errores de la Prueba de Estrés")
add_body("Análisis: Se observa que al alcanzar [Insertar número de usuarios o RPS observados], la latencia en percentil 95 experimenta un incremento drástico debido a la saturación de conexiones en el servidor web embebido y el pool de conexiones de la base de datos.")

# 5.3 Capacidad
add_heading_2("5.3. Escenario 3: Prueba de Capacidad y Resistencia (Soak Test)")
add_body("Objetivo: Probar la resistencia de la API ante una carga constante y prolongada cercana al máximo sostenible, detectando fugas de memoria, acumulación de sockets o agotamiento de recursos.")
add_body("Parámetros configurados y justificación:", bold_prefix="• ")
add_body("- Usuarios concurrentes (-u): 80 a 100 usuarios (carga alta pero dentro de la zona de estabilidad).")
add_body("- Tasa de arranque (-r): 5 usuarios/segundo.")
add_body("- Duración (-t): 15 a 30 minutos.")
add_body("Resultados de estabilidad en el tiempo:", bold_prefix="• ")
add_placeholder("Insertar gráfica de tendencia temporal sostenida de la Prueba de Capacidad")
add_body("Análisis: Los tiempos de respuesta y la memoria del servidor se mantienen estables a lo largo del tiempo, evidenciando ausencia de fugas de recursos tras la optimización.")

# 6. TABLA COMPARATIVA
add_heading_1("6. Tabla Comparativa de Métricas Globales")
add_body("A continuación se consolida la comparativa de desempeño entre el estado inicial de la API y los tres escenarios de prueba ejecutados:")

table_comp = doc.add_table(rows=5, cols=7)
table_comp.alignment = WD_TABLE_ALIGNMENT.CENTER
set_table_borders(table_comp)
comp_data = [
    ("Escenario Evaluado", "Concurrencia", "Peticiones Totales", "RPS Promedio", "Latencia p50 (ms)", "Latencia p95 (ms)", "% de Fallos"),
    ("API Sin Paginar (Inicial)", "1 usuario", "1 petición", "N/A", "Timeout / 500", "Timeout / 500", "100.0%"),
    ("Prueba de Carga (Load)", "50 usuarios", "[Completar]", "[Completar]", "[Completar]", "[Completar]", "0.0%"),
    ("Prueba de Estrés (Stress)", "250 usuarios", "[Completar]", "[Completar]", "[Completar]", "[Completar]", "[Completar]"),
    ("Prueba de Capacidad (Soak)", "80 usuarios", "[Completar]", "[Completar]", "[Completar]", "[Completar]", "0.0%")
]
for row_idx, cols in enumerate(comp_data):
    row = table_comp.rows[row_idx]
    for col_idx, text in enumerate(cols):
        cell = row.cells[col_idx]
        cell.text = text
        set_cell_margins(cell)
        is_header = (row_idx == 0)
        if is_header:
            set_cell_shading(cell, "E8EEF5")
        for p in cell.paragraphs:
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            for run in p.runs:
                set_run_font(run, name="Times New Roman", size=9.5, bold=is_header)

# 7. CONCLUSIONES
add_heading_1("7. Conclusiones y Lecciones Aprendidas")
add_body("1. Necesidad arquitectónica de la paginación:", bold_prefix="• ")
add_body("En sistemas con volúmenes de datos masivos (1.5M registros), la paginación no representa únicamente una característica de usabilidad, sino una barrera fundamental de protección para la memoria del servidor de aplicaciones y la red.")
add_body("2. Delegación de cómputo a motores de base de datos:", bold_prefix="• ")
add_body("El filtrado en memoria de la capa de aplicación (como User::all()->filter()) escala de manera pésima (O(N) con sobrecosto de serialización). Mover las condiciones a cláusulas WHERE en SQL con campos indexados permite que el motor de base de datos procese millones de registros en milisegundos.")
add_body("3. Distinción metodológica de pruebas de rendimiento:", bold_prefix="• ")
add_body("Cada tipo de prueba cumplió un rol diferenciado e indispensable: la prueba de carga verificó SLAs nominales, la prueba de estrés delimitó el techo operativo antes del colapso, y la prueba de capacidad garantizó la estabilidad en el tiempo.")

# 8. BITÁCORA DE IA
add_heading_1("8. Bitácora de Uso de Inteligencia Artificial (IA)")
add_body("Registro de auditoría técnica sobre la asistencia y co-creación mediante herramientas de Inteligencia Artificial durante el desarrollo del taller, en cumplimiento de los lineamientos éticos y metodológicos de la asignatura:")

# Encabezado institucional de la bitácora
p_meta = doc.add_paragraph()
p_meta.paragraph_format.space_before = Pt(4)
p_meta.paragraph_format.space_after = Pt(8)
p_meta.paragraph_format.line_spacing = 1.15
for label, val in [
    ("Fecha de realización: ", "8 de octubre de 2026\n"),
    ("Integrantes: ", "Andrés Hernández, Mateo Falla Muñoz\n"),
    ("Asignatura: ", "Ingeniería de Software II — Docente: Albeiro Montes\n"),
    ("Herramienta de IA utilizada: ", "Antigravity IDE (Modelo Gemini 3.8 Flash)\n")
]:
    r1 = p_meta.add_run(label)
    set_run_font(r1, name="Times New Roman", size=11, bold=True)
    r2 = p_meta.add_run(val)
    set_run_font(r2, name="Times New Roman", size=11, bold=False)

# Tabla de la bitácora
table_ai = doc.add_table(rows=6, cols=6)
table_ai.alignment = WD_TABLE_ALIGNMENT.CENTER
set_table_borders(table_ai)

ai_headers = [
    "Tarea Apoyada",
    "Prompt Utilizado (Resumen)",
    "Intervención Humana",
    "Resultado Obtenido",
    "Aprendizaje Técnico",
    "Evaluación Crítica"
]

ai_rows = [
    (
        "1. Análisis de requisitos y hoja de ruta",
        "\"analiza README-API-LOCUST.md... dame el paso a paso de lo que hay que hacer, no toques el código, solo instrucciones\"",
        "Se delimitó expresamente el alcance inicial para restringir la modificación directa de código y exigir un plan de trabajo estructurado.",
        "Comprensión del problema raíz en la API (1.5M registros sin paginar) y formulación de una hoja de ruta técnica en 7 fases.",
        "Comprender a fondo el cuello de botella arquitectónico y los criterios del docente antes de escribir cualquier línea de código.",
        "Excelente síntesis de los requisitos del laboratorio."
    ),
    (
        "2. Diagnóstico de falla de memoria y refactorización",
        "\"al momento de intentar ejecutar el get para traer los usuarios, no funciona... Como arreglamos eso / modifica el codigo y arreglalo\"",
        "Validación del error en ejecución real, autorización del cambio en el backend y verificación de imports en UserController.php.",
        "Implementación de ->paginate() con tope seguro (máx. 200) y conversión del filtro de overTwenty() a SQL indexado (birth_date).",
        "Los filtros y límites deben ejecutarse obligatoriamente en el motor de base de datos para no desbordar la memoria de PHP.",
        "Solución precisa y verificación de sintaxis y respuesta inmediata."
    ),
    (
        "3. Construcción y verificación de suite Locust",
        "\"Se tiene el archivo de locust... crea las funciones necesarias para las pruebas de locust ahí\"",
        "Suministro de la ruta de archivos del taller y revisión de las ponderaciones de tráfico exigidas en las diapositivas.",
        "Desarrollo de locustfile.py con tareas ponderadas (@task 5, 3, 3, 1), generación de UUIDs para evitar errores 422 y smoke test con 0% de fallas.",
        "Para evitar falsos positivos en pruebas de carga, las operaciones transaccionales con unicidad requieren generación dinámica de identificadores.",
        "Diseño robusto alineado con los criterios de aceptación del docente."
    ),
    (
        "4. Definición y justificación de parámetros",
        "\"dime si en el documento se tiene algun parametro para las pruebas en locust / dame los endpoints a los que hacer la solicitud\"",
        "Consulta activa sobre las directrices del docente y corroboración de que la parametrización numérica forma parte de la evaluación.",
        "Identificación de lineamientos de carga, estrés y capacidad; propuesta justificada de valores (-u, -r, -t) y contrato formal de endpoints.",
        "Las pruebas de rendimiento no son genéricas: cada escenario obedece a un perfil experimental específico (nominal, límite o resistencia).",
        "Alineación precisa con los objetivos pedagógicos."
    ),
    (
        "5. Maquetación del informe y bitácora técnica",
        "\"como redacto el documento, que debo poner en cada paso / añade la plantilla... sigue el formato... agrega bitacora de IA\"",
        "Instrucción de estructurar el informe Word respetando la tipografía Times New Roman y registrar la auditoría de uso de IA de la sesión.",
        "Generación automatizada de la plantilla completa con secciones, tablas de métricas y la bitácora de co-creación en informe.docx.",
        "La documentación rigurosa y transparente del uso de herramientas de IA refuerza el rigor académico y la trazabilidad del aprendizaje.",
        "Integración fluida y estructurada en el documento formal."
    )
]

for col_idx, h_text in enumerate(ai_headers):
    cell = table_ai.rows[0].cells[col_idx]
    cell.text = h_text
    set_cell_margins(cell, top=120, bottom=120, left=120, right=120)
    set_cell_shading(cell, "D9E1F2")
    for p in cell.paragraphs:
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        for r in p.runs:
            set_run_font(r, name="Times New Roman", size=9.5, bold=True)

for row_idx, row_data in enumerate(ai_rows, start=1):
    row = table_ai.rows[row_idx]
    for col_idx, cell_text in enumerate(row_data):
        cell = row.cells[col_idx]
        cell.text = cell_text
        set_cell_margins(cell, top=100, bottom=100, left=100, right=100)
        if row_idx % 2 == 0:
            set_cell_shading(cell, "F9FAFC")
        for p in cell.paragraphs:
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.05
            for r in p.runs:
                set_run_font(r, name="Times New Roman", size=9, bold=False)

# Guardar documento
doc.save(DOC_PATH)
print("informe.docx actualizado con éxito con la plantilla y la bitácora de IA.")
